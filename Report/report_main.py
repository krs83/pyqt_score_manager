import os
from pathlib import Path

import openpyxl

from Modules.const import ages_file

winners2 = {
    'name': [],
    'place': [],
    'ages': []
}


def log_report(msg):
    with open(r'D:\Python\scores_counter_bjj\Report\protocol\places.txt', 'a') as file_report:
        print(msg, file=file_report)


def form_winners(name_place1, age, place):
    winners2['name'].append(name_place1.strip())
    winners2['place'].append(place)
    winners2['ages'].append(age[:2])


def fill_places():
    for root, dirs, files in os.walk(Path(Path.cwd(), 'reports')):
        sum = 0
        medals1 = 0
        medals2 = 0
        medals3 = 0
        for i in files:
            wb = openpyxl.load_workbook(Path(root, i), read_only=True)
            w_sheets = wb.worksheets
            for j in range(len(w_sheets)):
                log_report(w_sheets[j]['D4'].value)
                name_place1 = w_sheets[j]['F23'].value
                name_place2 = w_sheets[j]['F24'].value
                name_place3 = w_sheets[j]['F25'].value
                age = w_sheets[j]['D4'].value

                # first place block
                log_report(f'I место: {name_place1}')
                form_winners(name_place1, age, 'I')
                medals1 += 1

                # second place block
                log_report(f'II место: {name_place2}')
                form_winners(name_place2, age, 'II')
                medals2 += 1

                # third place block
                if name_place3 is not None:
                    log_report(f'III место: {name_place3}')
                    form_winners(name_place3, age, 'III')
                    medals3 += 1
                log_report('')
            sum = sum + len(w_sheets)
            wb.close()
        log_report(f'Количество категорий {sum}')
        log_report(f'Медали I место в количестве {medals1} шт')
        log_report(f'Медали II место в количестве {medals2} шт')
        log_report(f'Медали III место в количестве {medals3} шт')
        log_report(f'Количество медалей {medals1 + medals2 + medals3}')


def check_empty_cell(wins, cell, w_sheet, i, j):
    if wins in cell.value:
        if w_sheet[f'H{7 + i}'].value:
            w_sheet[f'H{7 + i}'] = w_sheet[f'H{7 + i}'].value + f', {winners2["place"][j]}'
        else:
            w_sheet[f'H{7 + i}'] = winners2['place'][j]


def fill_protocol():
    file_path = Path(r'D:\Python\scores_counter_bjj\Report\protocol\Protocol_list.xlsx')
    wb = openpyxl.load_workbook(file_path)
    w_sheet = wb.worksheets[0]
    rows = len(list(w_sheet.rows))
    c = 0

    # fill up the  places
    for i, row in enumerate(w_sheet[f'C7:C{rows}']):
        for cell in row:
            for j, wins in enumerate(winners2['name']):
                if cell.value.strip() in winners2['name'][j]:
                    c += 1
                    check_empty_cell(wins, cell, w_sheet, i, j)
                    # print(c, cell.value, winners2['place'][j])

    wb.save(r'D:\Python\scores_counter_bjj\Report\protocol\Protocol_list.xlsx')


def fill_children_tournament():
    file_path = Path(r'D:\Python\scores_counter_bjj\Report\protocol\children.xlsx')
    wb = openpyxl.load_workbook(file_path)
    w_sheet = wb.worksheets[0]

    for i in range(len(winners2['name'])):
        num = i + 10
        w_sheet[f'C{num}'] = winners2['name'][i]
        w_sheet[f'D{num}'] = winners2['ages'][i]
        w_sheet[f'E{num}'] = winners2['place'][i]
        w_sheet[f'F{num}'] = winners2['place'][i]
        w_sheet[f'G{num}'] = 'грамота'

    wb.save(r'D:\Python\scores_counter_bjj\Report\protocol\children.xlsx')


def report_main():
    fill_places()
    fill_protocol()
    fill_children_tournament()


if __name__ == '__main__':
    report_main()
