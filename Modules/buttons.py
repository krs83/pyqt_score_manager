import os

import openpyxl
from PyQt5.QtWidgets import QWidget

from Modules.const import *
from Modules.flags import Flags
from Modules.utilities import logging

from UI.counter_screen import Ui_MainWindow


class TabButtons(QWidget):
    def __init__(self):
        super().__init__()
        self.ui = Ui_MainWindow()
        self.flags = Flags()

        self.ages_count = len(ages_file)
        self.sheet_names = ['10_', '9_', '8_', '7_', '6_', '5_', '4_', '3_', '2_']
        self.main_key = None

        # sample definitions
        self.file_path = f'{file_name}{sample}'
        self.wb = openpyxl.load_workbook(self.file_path)

        # list of buttons
        self.btns_age4 = []
        self.btns_age6 = []
        self.btns_age7 = []
        self.btns_age8 = []
        self.btns_age9 = []
        self.btns_age10 = []
        self.btns_age12 = []
        self.btns_age14 = []
        self.btns_age16 = []
        self.btns_age18 = []
        self.btns_age24 = []

        # list of age categories
        self.index_age4 = []
        self.index_age6 = []
        self.index_age7 = []
        self.index_age8 = []
        self.index_age9 = []
        self.index_age10 = []
        self.index_age12 = []
        self.index_age14 = []
        self.index_age16 = []
        self.index_age18 = []
        self.index_age24 = []

    def set_tab_buttons(self, age_file, age_btn_num, l_btns, l_index):
        if os.path.isfile(f'{file_name}\\{age_file}'):
            self.file_path = f'{file_name}{age_file}'
            self.wb = openpyxl.load_workbook(self.file_path)
            for c, i in enumerate(range(len(self.wb.sheetnames))):
                for j in self.sheet_names:
                    if self.wb.sheetnames[i].__contains__(j):
                        eval(f'self.ui.btn_{age_btn_num}_{c + 1}').setText(j + 'участников')
                        l_btns.append(f'self.ui.btn_{age_btn_num}_{c + 1}')

                        # формирует порядок таблиц участников в файле возрастной категории
                        if self.wb.sheetnames[i].__contains__('10_'):
                            l_index.append(self.wb.sheetnames[i][:2])
                        else:
                            l_index.append(self.wb.sheetnames[i][:1])

            # hide unusable buttons in tab
            for i in range(12):
                if eval(f'self.ui.btn_{age_btn_num}_{i + 1}').text() == '':
                    eval(f'self.ui.btn_{age_btn_num}_{i + 1}').setHidden(True)

    def handler_btns(self, age, l_btns, l_index):
        if os.path.isfile(f'{file_name}\\{age}'):
            for c, i in enumerate(l_btns):
                eval(i).clicked.connect(lambda state, index=c, btn=i: self.on_click_btns(index, btn, age, l_index))

    def spin_handlers(self):
        for i in range(2, 11):
            self.ui.btn_spin.clicked.connect(eval(f'self.spin_manage_{i}'))

    def on_click_btns(self, index, btn, age, l_index):
        self.main_key = str(f'{t}{l_index[index]}-' + age + str(f'{b}{index + 1}'))
        if os.path.isfile(f'{file_name}\\{age}'):
            eval(f'self.set_table_{l_index[index]}')(age=age, btn=eval(btn),
                                                     table=index, main_key=str(f'{t}{l_index[index]}-')
                                                     + str(age) + str(f'{b}{index + 1}'))

    #TODO: сократить код
    def backward(self):
        logging(cancel)
        self.on_off_spin(False)
        self.ui.btn_all_tables.setDisabled(False)
        for i in buttons4:
            eval(i).setStyleSheet(non_active_btn)
        for i in buttons6:
            eval(i).setStyleSheet(non_active_btn)
        for i in buttons7:
            eval(i).setStyleSheet(non_active_btn)
        for i in buttons8:
            eval(i).setStyleSheet(non_active_btn)
        for i in buttons9:
            eval(i).setStyleSheet(non_active_btn)
        for i in buttons10:
            eval(i).setStyleSheet(non_active_btn)
        for i in buttons12:
            eval(i).setStyleSheet(non_active_btn)
        for i in buttons14:
            eval(i).setStyleSheet(non_active_btn)
        for i in buttons16:
            eval(i).setStyleSheet(non_active_btn)
        for i in buttons18:
            eval(i).setStyleSheet(non_active_btn)
        for i in buttons24:
            eval(i).setStyleSheet(non_active_btn)

        for i in range(9):
            self.flags.is_table[i] = False

    def on_off_spin(self, boolean):
        self.ui.spinBox.setDisabled(boolean)
        self.ui.btn_spin.setDisabled(boolean)
        self.ui.laps_box.setDisabled(boolean)
